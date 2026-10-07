import os
import re
import json
import xml.etree.ElementTree as ET
import queue
import time
import threading
from boomi_migration_agent.boomi_client import BoomiClient

# Register namespaces for XML processing
ET.register_namespace('bns', 'http://api.platform.boomi.com/')

from concurrent.futures import ThreadPoolExecutor, as_completed

class MigrationTools:
    def __init__(self, client: BoomiClient = None):
        self.client = client or BoomiClient()
        self.folders_cache = {}
        self.root_folder_name = None
        self.root_folder_names = []

    def get_root_folder_name(self):
        """Returns dynamically detected root folder name for the current Boomi account."""
        if self.root_folder_name:
            return self.root_folder_name
        if not self.folders_cache:
            self.preload_folders()
        return self.root_folder_name or "Root"

    def _normalize_path(self, raw_path: str) -> str:
        """
        Dynamically normalizes folder paths under the account's actual root folder.
        Preserves existing multi-level paths and prevents hardcoding.
        """
        clean = (raw_path or "").replace("\\", "/").strip("/")
        root_name = self.get_root_folder_name()
        
        if not clean or clean == "Root":
            return root_name
            
        # Check if clean already starts with any identified root folder
        for r in (self.root_folder_names or [root_name]):
            if r and (clean == r or clean.startswith(f"{r}/")):
                return clean
                
        if root_name and root_name != "Root":
            if clean != root_name and not clean.startswith(f"{root_name}/"):
                return f"{root_name}/{clean}"
                
        return clean

    def preload_folders(self):
        """Preloads all folders across the account in one paginated query so all components have instant fullPath resolution."""
        if self.folders_cache:
            return self.folders_cache
        try:
            folders = self.client.query("Folder", {})

            # Step 1: Detect root folders dynamically from the account
            # In Boomi, root folders have empty/null parentId or parentId == ''
            root_candidates = [
                f for f in folders 
                if not (f.get("parentId") or "").strip() and not f.get("deleted", False)
            ]
            self.root_folder_names = [
                (f.get("name") or f.get("fullPath") or "").strip() 
                for f in root_candidates 
                if (f.get("name") or f.get("fullPath") or "").strip()
            ]
            if self.root_folder_names:
                self.root_folder_name = self.root_folder_names[0]
                
            # Fallback 1: Derive from top prefix of any existing folder's fullPath
            if not self.root_folder_name:
                for f in folders:
                    fp = (f.get("fullPath") or "").replace("\\", "/").strip("/")
                    if fp and not f.get("deleted", False):
                        prefix = fp.split("/")[0].strip()
                        if prefix and prefix not in self.root_folder_names:
                            self.root_folder_names.append(prefix)
                if self.root_folder_names:
                    self.root_folder_name = self.root_folder_names[0]
                    
            # Fallback 2: Default if account has no folders
            if not self.root_folder_name:
                self.root_folder_name = "Root"
                self.root_folder_names = ["Root"]

            for f in folders:
                fid = f.get("id")
                if fid:
                    fpath = f.get("fullPath") or f.get("name") or self.root_folder_name
                    self.folders_cache[fid] = self._normalize_path(fpath)
        except Exception as e:
            print(f"Warning: preload_folders error: {e}")
        return self.folders_cache

    def resolve_folders_for_items(self, items):
        """Resolves folder fullPath on-demand or from preloaded cache for the given components."""
        if not self.folders_cache:
            self.preload_folders()
            
        root_name = self.get_root_folder_name()
        for item in items:
            fid = item.get("folderId")
            fpath = self.folders_cache.get(fid) or item.get("folderName") or root_name
            item["folderPath"] = self._normalize_path(fpath)
            
        return self.folders_cache

    def get_folders_map(self):
        """Returns cached folders map or empty dict"""
        return self.folders_cache

    def scan_targeted_folder(self, folder_filter: str, progress_callback=None):
        """
        Targeted Fast Discovery Scan for a specific folder path:
        1. Queries Boomi Folder API to locate the specified folder.
        2. Recursively collects all child subfolders in the tree.
        3. Concurrently queries ComponentMetadata scoped strictly to those folder IDs.
        4. Traces where-used dependencies (Maps, Caches, DB Processes) for DB components in the folder.
        5. Returns scoped results in seconds.
        """
        raw_filter = (folder_filter or "").strip()
        clean_path = raw_filter.replace("\\", "/").strip("/")
        folder_name = clean_path.split("/")[-1].strip()

        if progress_callback:
            progress_callback("connections", f"Locating folder '{folder_name}'...", 1)

        # 1. Locate folder by exact name or substring
        matched_folders = []
        try:
            matched_folders = self.client.query("Folder", {"QueryFilter": {"expression": {
                "operator": "EQUALS", "property": "name", "argument": [folder_name]
            }}})
        except Exception as e:
            print(f"Warning: Folder query EQUALS error: {e}")

        if not matched_folders:
            try:
                matched_folders = self.client.query("Folder", {"QueryFilter": {"expression": {
                    "operator": "LIKE", "property": "name", "argument": [f"%{folder_name}%"]
                }}})
            except Exception as e:
                print(f"Warning: Folder query LIKE error: {e}")

        if not matched_folders:
            if progress_callback:
                progress_callback("connections", f"Folder '{raw_filter}' not found", 1)
            return {
                "folder_not_found": True,
                "searched_folder": raw_filter,
                "target_folder": folder_name,
                "connections": [],
                "operations": [],
                "profiles": [],
                "maps": [],
                "caches": [],
                "processes": []
            }

        # If multiple folders match, check if fullPath matches clean_path
        target_root = matched_folders[0]
        if len(matched_folders) > 1 and "/" in clean_path:
            for mf in matched_folders:
                fp = (mf.get("fullPath") or "").replace("\\", "/").strip("/")
                if fp.lower().endswith(clean_path.lower()):
                    target_root = mf
                    break

        target_folder_name = target_root.get("name") or folder_name

        # 2. Recursively gather all subfolder IDs and full paths
        all_target_folders = {target_root["id"]: target_root}
        root_path = target_root.get("fullPath") or target_root.get("name") or self.get_root_folder_name()
        root_path = self._normalize_path(root_path)
        self.folders_cache[target_root["id"]] = root_path

        curr_ids = [target_root["id"]]
        while curr_ids:
            next_ids = []
            for cid in curr_ids:
                try:
                    subs = self.client.query("Folder", {"QueryFilter": {"expression": {
                        "operator": "EQUALS", "property": "parentId", "argument": [cid]
                    }}})
                    for s in subs:
                        sid = s.get("id")
                        if sid and sid not in all_target_folders:
                            all_target_folders[sid] = s
                            fpath = s.get("fullPath") or s.get("name") or self.get_root_folder_name()
                            fpath = self._normalize_path(fpath)
                            self.folders_cache[sid] = fpath
                            next_ids.append(sid)
                except Exception as e:
                    print(f"Warning: Error traversing subfolders for {cid}: {e}")
            curr_ids = next_ids

        if progress_callback:
            progress_callback("connections", f"Scanning {len(all_target_folders)} folder(s)...", 1)

        # 3. Concurrently fetch all current active components in those folder IDs
        target_fids = list(all_target_folders.keys())
        all_folder_components = []

        def _fetch_folder_comps(fid):
            q = {"QueryFilter": {"expression": {"operator": "AND", "nestedExpression": [
                {"operator": "EQUALS", "property": "folderId", "argument": [fid]},
                {"operator": "EQUALS", "property": "currentVersion", "argument": ["true"]},
                {"operator": "EQUALS", "property": "deleted", "argument": ["false"]}
            ]}}}
            items = []
            try:
                items = self.client.query("ComponentMetadata", q)
            except Exception as e:
                print(f"Warning: Error fetching components for folder {fid}: {e}")

            if getattr(self.client, "branch_name", None):
                bq = {"QueryFilter": {"expression": {"operator": "AND", "nestedExpression": [
                    {"operator": "EQUALS", "property": "folderId", "argument": [fid]},
                    {"operator": "EQUALS", "property": "branchName", "argument": [self.client.branch_name]},
                    {"operator": "EQUALS", "property": "deleted", "argument": ["false"]}
                ]}}}
                try:
                    b_items = self.client.query("ComponentMetadata", bq)
                    item_map = {x["componentId"]: x for x in items}
                    for bi in b_items:
                        item_map[bi["componentId"]] = bi
                    items = list(item_map.values())
                except Exception as e:
                    print(f"Warning: Error fetching branch components for folder {fid}: {e}")

            return items

        with ThreadPoolExecutor(max_workers=min(25, max(len(target_fids), 1))) as exec:
            results_list = list(exec.map(_fetch_folder_comps, target_fids))
            for res in results_list:
                all_folder_components.extend(res)

        comp_by_id = {x["componentId"]: x for x in all_folder_components}

        # 4. Extract legacy DB components
        legacy_conns = [x for x in all_folder_components if x.get("type") == "connector-settings" and x.get("subType") == "database"]
        legacy_opers = [x for x in all_folder_components if x.get("type") == "connector-action" and x.get("subType") == "database"]
        legacy_profiles = [x for x in all_folder_components if x.get("type") == "profile.db"]

        if progress_callback:
            progress_callback("connections", len(legacy_conns), 1)
            progress_callback("operations", len(legacy_opers), 1)
            progress_callback("profiles", len(legacy_profiles), 1)
            progress_callback("maps_caches", "Resolving references...", 1)
            progress_callback("processes", "Resolving DB processes...", 1)

        # 5. Resolve references for Maps, Caches, and DB Processes IN PARALLEL
        folder_maps = {}
        folder_caches = {}
        folder_procs = {}

        def _resolve_profile_refs():
            prof_targets = [x["componentId"] for x in legacy_profiles]
            if prof_targets:
                pids = self._collect_parent_ids(prof_targets, max_workers=min(20, len(prof_targets)), progress_cb=progress_callback, cat_name="maps_caches")
                for pid in pids:
                    if pid in comp_by_id:
                        comp = comp_by_id[pid]
                        if comp.get("type") == "transform.map":
                            folder_maps[pid] = comp
                        elif comp.get("type") == "documentcache":
                            folder_caches[pid] = comp

        def _resolve_conn_oper_refs():
            conn_oper_targets = [x["componentId"] for x in legacy_conns] + [x["componentId"] for x in legacy_opers]
            if conn_oper_targets:
                pids = self._collect_parent_ids(conn_oper_targets, max_workers=min(20, len(conn_oper_targets)), progress_cb=progress_callback, cat_name="processes")
                for pid in pids:
                    if pid in comp_by_id and comp_by_id[pid].get("type") == "process":
                        folder_procs[pid] = comp_by_id[pid]

        with ThreadPoolExecutor(max_workers=2) as parallel_exec:
            f1 = parallel_exec.submit(_resolve_profile_refs)
            f2 = parallel_exec.submit(_resolve_conn_oper_refs)
            f1.result()
            f2.result()

        # Transitive processes from discovered maps and caches
        inter_ids = list(folder_maps.keys()) + list(folder_caches.keys())
        if inter_ids:
            inter_pids = self._collect_parent_ids(inter_ids, max_workers=min(20, len(inter_ids)))
            for pid in inter_pids:
                if pid in comp_by_id and comp_by_id[pid].get("type") == "process":
                    folder_procs[pid] = comp_by_id[pid]

        legacy_maps = list(folder_maps.values())
        legacy_caches = list(folder_caches.values())
        legacy_processes = list(folder_procs.values())

        # Enrich all items with folderPath
        root_name = self.get_root_folder_name()
        for item in (legacy_conns + legacy_opers + legacy_profiles + legacy_maps + legacy_caches + legacy_processes):
            fid = item.get("folderId")
            fpath = self.folders_cache.get(fid) or item.get("folderName") or root_name
            item["folderPath"] = self._normalize_path(fpath)

        if progress_callback:
            progress_callback("maps_caches", f"{len(legacy_maps)} Maps, {len(legacy_caches)} Caches", 1)
            progress_callback("processes", f"{len(legacy_processes)} DB Processes", 1)

        return {
            "target_folder": target_folder_name,
            "root_folder_name": root_name,
            "connections": legacy_conns,
            "operations": legacy_opers,
            "profiles": legacy_profiles,
            "maps": legacy_maps,
            "caches": legacy_caches,
            "processes": legacy_processes
        }

    def _get_parent_refs(self, c_id):
        max_retries = 4
        for attempt in range(max_retries):
            try:
                raw = self.client.query("ComponentReference", {"QueryFilter": {"expression": {
                    "operator": "EQUALS", "property": "componentId", "argument": [c_id]
                }}})
                pids = []
                for item in raw:
                    for ref in item.get("references", []):
                        pid = ref.get("parentComponentId")
                        if pid:
                            pids.append(pid)
                return pids
            except Exception as e:
                if attempt == max_retries - 1:
                    print(f"Warning: Failed _get_parent_refs for {c_id} after {max_retries} attempts: {e}")
                    return []
                time.sleep(0.5 * (2 ** attempt))

    def _collect_parent_ids(self, target_ids, max_workers=20, progress_cb=None, cat_name=""):
        parent_ids = set()
        id_list = list(set(target_ids))
        total = len(id_list)
        if not total:
            return parent_ids
        completed = 0
        with ThreadPoolExecutor(max_workers=min(max_workers, total)) as exec:
            futures = {exec.submit(self._get_parent_refs, cid): cid for cid in id_list}
            for f in as_completed(futures):
                try:
                    pids = f.result()
                    for pid in pids:
                        parent_ids.add(pid)
                except Exception:
                    pass
                completed += 1
                if progress_cb and (completed % 100 == 0 or completed == total):
                    progress_cb(cat_name, f"Resolving references ({completed}/{total})...", completed)
        return parent_ids

    def _batch_resolve_metadata(self, comp_ids, batch_size=25):
        if not comp_ids:
            return {}
        id_list = list(set(comp_ids))
        resolved = {}

        def _resolve_chunk(chunk):
            q = {
                "QueryFilter": {
                    "expression": {
                        "operator": "and",
                        "nestedExpression": [
                            {"operator": "or", "nestedExpression": [
                                {"operator": "EQUALS", "property": "componentId", "argument": [cid]} for cid in chunk
                            ]},
                            {"operator": "EQUALS", "property": "deleted", "argument": ["false"]}
                        ]
                    }
                }
            }
            for attempt in range(3):
                try:
                    items = self.client.query("ComponentMetadata", q)
                    return items
                except Exception as e:
                    if attempt == 2:
                        break
                    time.sleep(0.5 * (attempt + 1))

            fallback_items = []
            for cid in chunk:
                for attempt in range(3):
                    try:
                        raw = self.client.query("ComponentMetadata", {"QueryFilter": {"expression": {
                            "operator": "and", "nestedExpression": [
                                {"operator": "EQUALS", "property": "componentId", "argument": [cid]},
                                {"operator": "EQUALS", "property": "deleted", "argument": ["false"]}
                            ]
                        }}})
                        if raw:
                            fallback_items.extend(raw)
                        break
                    except Exception:
                        time.sleep(0.3 * (attempt + 1))
            return fallback_items

        chunks = [id_list[i:i + batch_size] for i in range(0, len(id_list), batch_size)]
        with ThreadPoolExecutor(max_workers=min(8, max(len(chunks), 1))) as exec:
            chunk_results = list(exec.map(_resolve_chunk, chunks))
            
            by_cid = {}
            for items in chunk_results:
                for it in items:
                    cid = it.get("componentId")
                    if cid:
                        by_cid.setdefault(cid, []).append(it)

            for cid, items in by_cid.items():
                best = None
                if getattr(self.client, "branch_name", None):
                    best = next((x for x in items if x.get("branchName") == self.client.branch_name and x.get("deleted") is False), None)
                if not best:
                    best = next((x for x in items if x.get("currentVersion") is True and x.get("deleted") is False), None)
                if not best:
                    best = next((x for x in items if x.get("deleted") is False), None) or items[0]
                
                fid = best.get("folderId")
                fpath = self.folders_cache.get(fid) or best.get("folderName") or self.get_root_folder_name()
                best["folderPath"] = self._normalize_path(fpath)
                resolved[cid] = best

        return resolved

    def _discover_maps_and_caches(self, legacy_profiles, progress_callback=None):
        prof_ids = [p["componentId"] for p in legacy_profiles]
        if progress_callback:
            progress_callback("maps_caches", f"Analyzing {len(prof_ids)} profiles...", 1)

        parent_ids = self._collect_parent_ids(prof_ids, max_workers=16, progress_cb=progress_callback, cat_name="maps_caches")

        if progress_callback:
            progress_callback("maps_caches", f"Resolving metadata ({len(parent_ids)} items)...", 1)

        metas = self._batch_resolve_metadata(list(parent_ids))
        maps = [m for m in metas.values() if m.get("type") == "transform.map"]
        caches = [c for c in metas.values() if c.get("type") == "documentcache"]

        if progress_callback:
            progress_callback("maps_caches", f"{len(maps)} Maps, {len(caches)} Caches", 1)

        return maps, caches

    def _discover_db_processes(self, legacy_conns, legacy_opers, progress_callback=None):
        target_ids = [c["componentId"] for c in legacy_conns] + [o["componentId"] for o in legacy_opers]
        if progress_callback:
            progress_callback("processes", f"Analyzing {len(target_ids)} DB components...", 1)

        parent_ids = self._collect_parent_ids(target_ids, max_workers=16, progress_cb=progress_callback, cat_name="processes")

        if progress_callback:
            progress_callback("processes", f"Resolving metadata ({len(parent_ids)} items)...", 1)

        metas = self._batch_resolve_metadata(list(parent_ids))
        processes = [p for p in metas.values() if p.get("type") == "process"]

        if progress_callback:
            progress_callback("processes", f"{len(processes)} DB Processes", 1)

        return processes

    def scan_legacy_components(self, folder_filter=None, progress_callback=None):
        """
        High-Performance Concurrent Discovery Scan:
        1. If folder_filter is provided, runs targeted fast discovery scoped to that folder tree.
        2. Otherwise, preloads all account folders and fetches legacy DB Connections, Operations, and Profiles account-wide in parallel.
        3. Concurrently runs Maps/Caches discovery and DB Process discovery in PARALLEL.
        4. Queries transitive DB processes from discovered Maps and Caches.
        5. Assigns exact hierarchical full folder paths under the dynamic account root folder.
        """
        raw_filter = (folder_filter or "").strip()
        if raw_filter:
            return self.scan_targeted_folder(raw_filter, progress_callback=progress_callback)

        # Notify UI immediately that all streams are actively running
        if progress_callback:
            progress_callback("connections", 0, 1)
            progress_callback("operations", 0, 1)
            progress_callback("profiles", 0, 1)
            progress_callback("maps_caches", "Preparing map discovery...", 1)
            progress_callback("processes", "Preparing DB process discovery...", 1)

        def build_query_payload(comp_type, sub_type=None):
            nested = [
                {"operator": "EQUALS", "property": "type", "argument": [comp_type]},
                {"operator": "EQUALS", "property": "currentVersion", "argument": ["true"]},
                {"operator": "EQUALS", "property": "deleted", "argument": ["false"]}
            ]
            if sub_type:
                nested.append({"operator": "EQUALS", "property": "subType", "argument": [sub_type]})
            return {"QueryFilter": {"expression": {"operator": "and", "nestedExpression": nested}}}

        def _fetch_by_type(comp_type, sub_type=None, cat_name=""):
            cb = (lambda count, page: progress_callback(cat_name, count, page)) if progress_callback else None
            payload = build_query_payload(comp_type, sub_type)
            items = self.client.query("ComponentMetadata", payload, progress_callback=cb)

            if getattr(self.client, "branch_name", None):
                branch_nested = [
                    {"operator": "EQUALS", "property": "type", "argument": [comp_type]},
                    {"operator": "EQUALS", "property": "branchName", "argument": [self.client.branch_name]},
                    {"operator": "EQUALS", "property": "deleted", "argument": ["false"]}
                ]
                if sub_type:
                    branch_nested.append({"operator": "EQUALS", "property": "subType", "argument": [sub_type]})
                try:
                    branch_items = self.client.query("ComponentMetadata", {"QueryFilter": {"expression": {"operator": "and", "nestedExpression": branch_nested}}})
                    item_map = {x["componentId"]: x for x in items}
                    for bi in branch_items:
                        item_map[bi["componentId"]] = bi
                    items = list(item_map.values())
                except Exception as e:
                    print(f"Warning: Error fetching branch items for {comp_type}: {e}")

            return items

        # Phase 1: Concurrently fetch all base components and preload folders
        with ThreadPoolExecutor(max_workers=4) as phase1_exec:
            f_conns = phase1_exec.submit(_fetch_by_type, "connector-settings", "database", "connections")
            f_opers = phase1_exec.submit(_fetch_by_type, "connector-action", "database", "operations")
            f_profs = phase1_exec.submit(_fetch_by_type, "profile.db", None, "profiles")
            f_folders = phase1_exec.submit(self.preload_folders)

            legacy_conns = f_conns.result()
            legacy_opers = f_opers.result()
            legacy_profiles = f_profs.result()
            f_folders.result()

        if progress_callback:
            progress_callback("connections", len(legacy_conns), 1)
            progress_callback("operations", len(legacy_opers), 1)
            progress_callback("profiles", len(legacy_profiles), 1)

        # Phase 2: Concurrently discover Maps/Caches AND DB Processes IN PARALLEL
        with ThreadPoolExecutor(max_workers=2) as phase2_exec:
            f_maps_caches = phase2_exec.submit(self._discover_maps_and_caches, legacy_profiles, progress_callback)
            f_processes = phase2_exec.submit(self._discover_db_processes, legacy_conns, legacy_opers, progress_callback)

            legacy_maps, legacy_caches = f_maps_caches.result()
            initial_processes = f_processes.result()

        # Phase 3: Transitive DB Processes referencing discovered Maps & Caches
        procs_dict = {p["componentId"]: p for p in initial_processes}
        inter_ids = [m["componentId"] for m in legacy_maps] + [c["componentId"] for c in legacy_caches]
        if inter_ids:
            if progress_callback:
                progress_callback("processes", f"Checking transitive usage for {len(inter_ids)} Maps/Caches...", 1)
            inter_parent_ids = self._collect_parent_ids(inter_ids, max_workers=16, progress_cb=progress_callback, cat_name="processes")
            unresolved_proc_ids = [pid for pid in inter_parent_ids if pid not in procs_dict]
            if unresolved_proc_ids:
                new_metas = self._batch_resolve_metadata(unresolved_proc_ids)
                for m in new_metas.values():
                    if m.get("type") == "process":
                        procs_dict[m["componentId"]] = m

        legacy_processes = list(procs_dict.values())

        # Assign full hierarchical folders
        self.resolve_folders_for_items(legacy_conns + legacy_opers + legacy_profiles + legacy_maps + legacy_caches + legacy_processes)

        # Enrich all items with folderPath anchored under dynamic root folder
        root_name = self.get_root_folder_name()
        for item in (legacy_conns + legacy_opers + legacy_profiles + legacy_maps + legacy_caches + legacy_processes):
            fid = item.get("folderId")
            fpath = item.get("folderPath") or self.folders_cache.get(fid) or item.get("folderName") or root_name
            item["folderPath"] = self._normalize_path(fpath)

        if progress_callback:
            progress_callback("maps_caches", f"{len(legacy_maps)} Maps, {len(legacy_caches)} Caches", 1)
            progress_callback("processes", f"{len(legacy_processes)} DB Processes", 1)

        return {
            "root_folder_name": root_name,
            "connections": legacy_conns,
            "operations": legacy_opers,
            "profiles": legacy_profiles,
            "maps": legacy_maps,
            "caches": legacy_caches,
            "processes": legacy_processes
        }

    def find_where_used(self, component_id, version):
        """
        Step 2 Where-Used: Queries component references to find all dependencies.
        """
        # referenced-by (other components that point to this one)
        ref_by_raw = []
        try:
            ref_by_raw = self.client.query("ComponentReference", {"QueryFilter": {"expression": {
                "operator": "EQUALS", "property": "componentId", "argument": [component_id]
            }}})
        except Exception as e:
            print(f"Warning: Failed to fetch referenced-by for {component_id}: {e}")

        ref_by = []
        for item in ref_by_raw:
            ref_by.extend(item.get("references", []))

        # references (components that this one points to)
        ref_to_raw = []
        try:
            ref_to_raw = self.client.query("ComponentReference", {"QueryFilter": {"expression": {
                "operator": "and", "nestedExpression": [
                    {"operator": "EQUALS", "property": "parentComponentId", "argument": [component_id]},
                    {"operator": "EQUALS", "property": "parentVersion", "argument": [str(version)]}
                ]
            }}})
        except Exception as e:
            print(f"Warning: Failed to fetch references for {component_id} version {version}: {e}")
            
        ref_to = []
        for item in ref_to_raw:
            ref_to.extend(item.get("references", []))
            
        return {
            "referenced_by": ref_by,
            "references": ref_to
        }

    def resolve_component_details(self, component_id):
        if not component_id:
            return None
        if not hasattr(self, 'metadata_cache'):
            self.metadata_cache = {}
        if component_id in self.metadata_cache:
            return self.metadata_cache[component_id]
            
        try:
            results = self.client.query("ComponentMetadata", {"QueryFilter": {"expression": {
                "operator": "EQUALS", "property": "componentId", "argument": [component_id]
            }}})
            if results:
                meta = results[0]
                fid = meta.get("folderId")
                folder_path = self.get_folders_map().get(fid) if fid else None
                if not folder_path:
                    folder_path = meta.get("folderFullPath") or meta.get("folderName", "Unknown Folder")
                self.metadata_cache[component_id] = {
                    "name": meta.get("name"),
                    "type": meta.get("type"),
                    "folderName": meta.get("folderName", "Unknown Folder"),
                    "folderPath": folder_path
                }
                return self.metadata_cache[component_id]
        except Exception as e:
            print(f"Warning: Failed to resolve metadata for {component_id}: {e}")
            
        return None

    def generate_audit_report_md(self, scan_results, include_where_used=True, max_where_used=50):
        """
        Generates a detailed audit report of legacy database usage in Markdown format.
        """
        md = []
        md.append("# Boomi Legacy Database Component Audit Report")
        md.append("\nThis report lists all Database (Legacy) connections, operations, profiles, and their dependent workflows.\n")
        
        # 1. Connections Table
        md.append("## 1. Legacy Database Connections")
        if not scan_results["connections"]:
            md.append("*No legacy database connections found.*")
        else:
            md.append(f"**Total Found**: {len(scan_results['connections'])}\n")
            md.append("| Connection Name | Component ID | Folder Path | Version |")
            md.append("| --- | --- | --- | --- |")
            for c in scan_results["connections"][:200]:
                md.append(f"| {c['name']} | `{c['componentId']}` | {c['folderPath']} | {c['version']} |")
            if len(scan_results["connections"]) > 200:
                md.append(f"| ... and {len(scan_results['connections']) - 200} more | ... | ... | ... |")
        md.append("\n")

        # 2. Operations Table
        md.append("## 2. Legacy Database Operations")
        if not scan_results["operations"]:
            md.append("*No legacy database operations found.*")
        else:
            md.append(f"**Total Found**: {len(scan_results['operations'])}\n")
            md.append("| Operation Name | Component ID | Folder Path | Version |")
            md.append("| --- | --- | --- | --- |")
            for o in scan_results["operations"][:200]:
                md.append(f"| {o['name']} | `{o['componentId']}` | {o['folderPath']} | {o['version']} |")
            if len(scan_results["operations"]) > 200:
                md.append(f"| ... and {len(scan_results['operations']) - 200} more | ... | ... | ... |")
        md.append("\n")

        # 3. Profiles Table
        md.append("## 3. Legacy Database Profiles")
        if not scan_results["profiles"]:
            md.append("*No legacy database profiles found.*")
        else:
            md.append(f"**Total Found**: {len(scan_results['profiles'])}\n")
            md.append("| Profile Name | Component ID | Folder Path | Version |")
            md.append("| --- | --- | --- | --- |")
            for p in scan_results["profiles"][:200]:
                md.append(f"| {p['name']} | `{p['componentId']}` | {p['folderPath']} | {p['version']} |")
            if len(scan_results["profiles"]) > 200:
                md.append(f"| ... and {len(scan_results['profiles']) - 200} more | ... | ... | ... |")
        md.append("\n")

        # 4. Detailed Dependency Mapping (Where-Used)
        if include_where_used:
            md.append("## 4. Where Used / Dependency Mapping")
            all_items = (
                [("Connection", c) for c in scan_results["connections"]] +
                [("Operation", o) for o in scan_results["operations"]] +
                [("Profile", p) for p in scan_results["profiles"]]
            )
            
            items_to_audit = all_items[:max_where_used]
            if len(all_items) > max_where_used:
                md.append(f"> [!NOTE]\n> Auditing dependencies for first {max_where_used} of {len(all_items)} total components. Use targeted folder filtering for specific integrations.\n")
                
            for comp_type, item in items_to_audit:
                comp_id = item["componentId"]
                name = item["name"]
                version = item["version"]
                
                md.append(f"### {comp_type}: {name} (`{comp_id}`)")
                
                refs = self.find_where_used(comp_id, version)
                ref_by = refs["referenced_by"]
                
                if not ref_by:
                    md.append("*   **Used By**: This component is not referenced anywhere (orphan).")
                else:
                    md.append("*   **Used By Dependencies**:")
                    for ref in ref_by:
                        parent_id = ref.get("parentComponentId")
                        parent_details = self.resolve_component_details(parent_id)
                        if parent_details:
                            parent_name = parent_details["name"]
                            parent_type = parent_details["type"]
                            parent_folder = parent_details["folderName"]
                            md.append(f"    *   **{parent_name}** (Type: `{parent_type}`, ID: `{parent_id}`, Folder: `{parent_folder}`)")
                        else:
                            md.append(f"    *   **Unnamed Parent** (ID: `{parent_id}`)")
                
                md.append("") # Empty line
                
        return "\n".join(md)


    def assemble_dbv2_connection_xml(self, legacy_xml_str, folder_id, new_name):
        """
        Translates Legacy Database Connection XML parameters to DB v2 Connection XML.
        Rule 1: Don't keep passwords; password fields are omitted or left blank.
        """
        root = ET.fromstring(legacy_xml_str)
        
        # Parse attributes from legacy root or DatabaseConnectionSettings
        settings = root.find('.//DatabaseConnectionSettings')
        if settings is None:
            raise ValueError("Invalid Database (Legacy) connection XML: DatabaseConnectionSettings element missing.")
            
        driver_id = settings.get("driverId", "")
        class_name = settings.get("className", "")
        host = settings.get("host", "")
        port = settings.get("port", "")
        dbname = settings.get("dbname", "")
        additional = settings.get("additional", "")
        username = settings.get("username", "")
        
        # Build JDBC URL
        if driver_id == "sqlserver":
            url = f"jdbc:sqlserver://{host}:{port};databaseName={dbname}{additional}"
        elif driver_id == "mysql":
            url = f"jdbc:mysql://{host}:{port}/{dbname}{additional}"
        elif driver_id == "oracle":
            url = f"jdbc:oracle:thin:@{host}:{port}:{dbname}{additional}"
        else:
            # Fallback/Custom Driver
            url = f"jdbc:generic://{host}:{port}/{dbname}{additional}"
            
        # Create DB V2 Connection XML structure
        # SubType for DB V2 is 'officialboomi-X3979C-dbv2da-prod'
        ET.register_namespace('bns', 'http://api.platform.boomi.com/')
        comp = ET.Element('{http://api.platform.boomi.com/}Component', {
            'componentId': '',
            'name': new_name,
            'type': 'connector-settings',
            'subType': 'officialboomi-X3979C-dbv2da-prod',
            'folderId': folder_id
        })
        
        enc_vals = ET.SubElement(comp, '{http://api.platform.boomi.com/}encryptedValues')
        # Empty encryptedValue element for password configuration (omitting value as per Rule 1)
        ET.SubElement(enc_vals, '{http://api.platform.boomi.com/}encryptedValue', {
            'isSet': 'false',
            'path': "//GenericConnectionConfig/field[@type='password']"
        })
        
        obj = ET.SubElement(comp, '{http://api.platform.boomi.com/}object')
        conn_config = ET.SubElement(obj, 'GenericConnectionConfig')
        
        # Add connection fields
        ET.SubElement(conn_config, 'field', {'id': 'url', 'type': 'string', 'value': url})
        ET.SubElement(conn_config, 'field', {'id': 'className', 'type': 'string', 'value': class_name})
        ET.SubElement(conn_config, 'field', {'id': 'username', 'type': 'string', 'value': username})
        ET.SubElement(conn_config, 'field', {'id': 'password', 'type': 'password'}) # Password is blank
        
        # Pooling fields (default to false/disabled)
        ET.SubElement(conn_config, 'field', {'id': 'enablePooling', 'type': 'boolean', 'value': 'false'})
        
        return ET.tostring(comp, encoding='utf-8').decode('utf-8')

    def parse_profile_db_sql(self, profile_xml_str):
        """
        Parses legacy profile.db XML to extract SQL statement, tableName, executionType,
        and database parameters/fields.
        """
        root = ET.fromstring(profile_xml_str)
        db_profile = root.find('.//DatabaseProfile')
        if db_profile is None:
            db_profile = root.find('.//{http://api.platform.boomi.com/}DatabaseProfile')
            if db_profile is None:
                raise ValueError("profile.db XML does not contain DatabaseProfile element.")
                
        stmt = db_profile.find('.//DBStatement')
        if stmt is None:
            raise ValueError("profile.db does not contain a DBStatement node.")
            
        sql_el = stmt.find('sql')
        sql_text = sql_el.text if sql_el is not None else ""
        
        fields = []
        for el in stmt.findall('.//DBFields/DatabaseElement'):
            fields.append({
                "key": el.get("key"),
                "name": el.get("name"),
                "dataType": el.get("dataType", "character")
            })
            
        params = []
        for el in stmt.findall('.//DBParameters/DatabaseElement'):
            params.append({
                "key": el.get("key"),
                "name": el.get("name"),
                "dataType": el.get("dataType", "character")
            })
            
        # Also check DBConditions in legacy profiles
        condition_elements = stmt.findall('.//DBCondition') + stmt.findall('.//DBConditions/DatabaseElement') + stmt.findall('.//Conditions/DatabaseElement')
        conditions_info = []
        for el in condition_elements:
            c_name = el.get("name")
            c_op = el.get("conditionOperator", "equal").lower()
            if c_name:
                conditions_info.append({
                    "name": c_name,
                    "operator": c_op,
                    "dataType": el.get("dataType", "character")
                })
                if not any(p.get("name") == c_name for p in params):
                    params.append({
                        "key": el.get("key"),
                        "name": c_name,
                        "dataType": el.get("dataType", "character")
                    })
            
        stmt_type = stmt.get("statementType", "select")
        table_name = stmt.get("tableName", "")
        
        # Operator mapping for SQL conditions
        op_map = {
            "equal": "= ?",
            "notequal": "<> ?",
            "greaterthan": "> ?",
            "lessthan": "< ?",
            "greaterthanorequal": ">= ?",
            "lessthanorequal": "<= ?",
            "like": "LIKE ?",
            "isnull": "IS NULL",
            "isnotnull": "IS NOT NULL"
        }
        
        # Synthesize SQL query ONLY for dynamicupdate based on profile logic
        # (dynamicinsert and dynamicdelete are handled natively by DBv2 without synthesized SQL)
        if not sql_text and table_name and stmt_type == "dynamicupdate" and fields:
            set_clause = ", ".join([f"{f['name']} = ?" for f in fields])
            if conditions_info:
                where_clauses = [f"{c['name']} {op_map.get(c['operator'], '= ?')}" for c in conditions_info]
                where_clause = " AND ".join(where_clauses)
            else:
                where_clause = " AND ".join([f"{f['name']} = ?" for f in fields])
            sql_text = f"UPDATE {table_name} SET {set_clause}" + (f" WHERE {where_clause}" if where_clause else "")

        return {
            "sql": sql_text,
            "statementType": stmt_type,
            "tableName": table_name,
            "fields": fields,
            "params": params
        }

    def generate_json_profile_xml(self, profile_name, fields, folder_id):
        """
        Generates a flat JSON profile XML representation based on database fields list
        """
        ET.register_namespace('bns', 'http://api.platform.boomi.com/')
        comp = ET.Element('{http://api.platform.boomi.com/}Component', {
            'componentId': '',
            'name': profile_name,
            'type': 'profile.json',
            'folderId': folder_id
        })
        
        ET.SubElement(comp, '{http://api.platform.boomi.com/}encryptedValues')
        obj = ET.SubElement(comp, '{http://api.platform.boomi.com/}object')
        
        json_profile = ET.SubElement(obj, 'JSONProfile', {'strict': 'false'})
        data_elements = ET.SubElement(json_profile, 'DataElements')
        
        root_val = ET.SubElement(data_elements, 'JSONRootValue', {
            'dataType': 'character',
            'isMappable': 'true',
            'isNode': 'true',
            'key': '1',
            'name': 'Root'
        })
        
        df_root = ET.SubElement(root_val, 'DataFormat')
        ET.SubElement(df_root, 'ProfileCharacterFormat')
        
        json_obj = ET.SubElement(root_val, 'JSONObject', {
            'isMappable': 'false',
            'isNode': 'true',
            'key': '2',
            'name': 'Object'
        })
        
        curr_key = 3
        for field in fields:
            db_type = field.get("dataType", "character").lower()
            if "integer" in db_type or "number" in db_type or "numeric" in db_type or "float" in db_type or "double" in db_type:
                js_type = "number"
            elif "boolean" in db_type:
                js_type = "boolean"
            elif "date" in db_type or "time" in db_type:
                js_type = "datetime"
            else:
                js_type = "character"
                
            entry = ET.SubElement(json_obj, 'JSONObjectEntry', {
                'dataType': js_type,
                'isMappable': 'true',
                'isNode': 'true',
                'key': str(curr_key),
                'name': field["name"]
            })
            
            df_entry = ET.SubElement(entry, 'DataFormat')
            if js_type == "number":
                ET.SubElement(df_entry, 'ProfileNumberFormat', {'numberFormat': ''})
            elif js_type == "boolean":
                ET.SubElement(df_entry, 'ProfileBooleanFormat')
            elif js_type == "datetime":
                ET.SubElement(df_entry, 'ProfileDateFormat')
            else:
                ET.SubElement(df_entry, 'ProfileCharacterFormat')
                
            curr_key += 1
            
        quals = ET.SubElement(root_val, 'Qualifiers')
        ET.SubElement(quals, 'QualifierList')
        
        ET.SubElement(json_profile, 'tagLists')
        
        return ET.tostring(comp, encoding='utf-8').decode('utf-8')

    def generate_dbv2_standard_response_profile(self, table_name, folder_id, action="UPDATE"):
        """
        Generates standard Database V2 Response JSON Profile.
        e.g. Database V2 {table_name} (TABLE) {action} Response
        Contains fields: Query, Rows Effected, Status.
        """
        clean_table = (table_name or "table").strip()
        resp_name = f"Database V2 {clean_table} (TABLE) {action} Response"
        fields = [
            {"name": "Query", "dataType": "character"},
            {"name": "Rows Effected", "dataType": "number"},
            {"name": "Status", "dataType": "character"}
        ]
        return resp_name, self.generate_json_profile_xml(resp_name, fields, folder_id)

    def assemble_dbv2_operation_xml(self, new_oper_name, connection_v2_id, sql_query, statement_type, request_profile_id, response_profile_id, folder_id, table_name=None, schema_name=""):
        """
        Assembles Database V2 Operation component XML.
        """
        ET.register_namespace('bns', 'http://api.platform.boomi.com/')
        comp = ET.Element('{http://api.platform.boomi.com/}Component', {
            'componentId': '',
            'name': new_oper_name,
            'type': 'connector-action',
            'subType': 'officialboomi-X3979C-dbv2da-prod',
            'folderId': folder_id
        })
        
        ET.SubElement(comp, '{http://api.platform.boomi.com/}encryptedValues')
        obj = ET.SubElement(comp, '{http://api.platform.boomi.com/}object')
        
        op_type = "GET"
        if statement_type in ["standardinsertupdatedelete", "dynamicinsert"]:
            op_type = "CREATE"
        elif statement_type in ["dynamicupdate"]:
            op_type = "UPDATE"
        elif statement_type in ["dynamicdelete"]:
            op_type = "DELETE"
            
        oper = ET.SubElement(obj, 'Operation', {
            'returnApplicationErrors': 'true',
            'trackResponse': 'false'
        })
        
        ET.SubElement(oper, 'Archiving', {'directory': '', 'enabled': 'false'})
        config = ET.SubElement(oper, 'Configuration')
        
        clean_table = (table_name or "table").strip()
        obj_type_id = clean_table.lower()
        obj_type_name = f"{clean_table} (TABLE)"

        gen_config_attrs = {
            'operationType': "EXECUTE" if op_type == "GET" else op_type,
            'requestProfile': request_profile_id or "",
            'requestProfileType': "json" if request_profile_id else "",
            'responseProfile': response_profile_id or "",
            'responseProfileType': "json" if response_profile_id else ""
        }
        
        # In DB V2:
        # customOperationType is used for GET (EXECUTE) and CREATE.
        # For UPDATE, customOperationType MUST NOT be set (causes "The 'UPDATE' action is no longer available" warning).
        if op_type == "GET":
            gen_config_attrs['customOperationType'] = "GET"
        elif op_type == "CREATE":
            gen_config_attrs['customOperationType'] = "CREATE"
            gen_config_attrs['objectTypeId'] = obj_type_id
            gen_config_attrs['objectTypeName'] = obj_type_name
        elif op_type == "UPDATE":
            gen_config_attrs['objectTypeId'] = obj_type_id
            gen_config_attrs['objectTypeName'] = obj_type_name
        elif op_type == "DELETE":
            gen_config_attrs['objectTypeId'] = obj_type_id
            gen_config_attrs['objectTypeName'] = obj_type_name

        gen_config = ET.SubElement(config, 'GenericOperationConfig', gen_config_attrs)
        
        if op_type == "GET":
            ET.SubElement(gen_config, 'field', {'id': 'GetType', 'type': 'string', 'value': 'Standard Get'})
            ET.SubElement(gen_config, 'field', {'id': 'INClause', 'type': 'boolean', 'value': 'false'})
            ET.SubElement(gen_config, 'field', {'id': 'query', 'type': 'string', 'value': sql_query or ""})
        elif op_type == "CREATE":
            insert_type = "Standard Insert" if sql_query else "Dynamic Insert"
            ET.SubElement(gen_config, 'field', {'id': 'InsertionType', 'type': 'string', 'value': insert_type})
            ET.SubElement(gen_config, 'field', {'id': 'schemaName', 'type': 'string', 'value': schema_name or ''})
            ET.SubElement(gen_config, 'field', {'id': 'query', 'type': 'string', 'value': sql_query or ""})
            ET.SubElement(gen_config, 'field', {'id': 'joinTransaction', 'type': 'boolean', 'value': 'false'})
            ET.SubElement(gen_config, 'field', {'id': 'CommitOption', 'type': 'string', 'value': 'Commit By Profile'})
            ET.SubElement(gen_config, 'field', {'id': 'batchCount', 'type': 'integer'})
            
            opts = ET.SubElement(gen_config, 'Options')
            qopts = ET.SubElement(opts, 'QueryOptions')
            fields_el = ET.SubElement(qopts, 'Fields')
            conn_obj = ET.SubElement(fields_el, 'ConnectorObject', {'name': obj_type_name})
            field_list = ET.SubElement(conn_obj, 'FieldList')
            ET.SubElement(field_list, 'ConnectorField', {'filterable': 'true', 'name': 'Query', 'selectable': 'true', 'selected': 'true', 'sortable': 'true'})
            ET.SubElement(field_list, 'ConnectorField', {'filterable': 'true', 'name': 'Rows Effected', 'selectable': 'true', 'selected': 'true', 'sortable': 'true'})
            ET.SubElement(field_list, 'ConnectorField', {'filterable': 'true', 'name': 'Status', 'selectable': 'true', 'selected': 'true', 'sortable': 'true'})
            ET.SubElement(qopts, 'Inputs')
        elif op_type == "UPDATE":
            ET.SubElement(gen_config, 'field', {'id': 'Type', 'type': 'string', 'value': 'Dynamic Update'})
            ET.SubElement(gen_config, 'field', {'id': 'schemaName', 'type': 'string', 'value': schema_name or ''})
            ET.SubElement(gen_config, 'field', {'id': 'query', 'type': 'string', 'value': sql_query or ""})
            ET.SubElement(gen_config, 'field', {'id': 'joinTransaction', 'type': 'boolean', 'value': 'false'})
            ET.SubElement(gen_config, 'field', {'id': 'CommitOption', 'type': 'string', 'value': 'Commit By Profile'})
            ET.SubElement(gen_config, 'field', {'id': 'batchCount', 'type': 'integer'})
            
            dyn_field = ET.SubElement(gen_config, 'dynamicOperationField', {
                'displayType': 'textarea',
                'id': 'query',
                'label': 'SQL Query',
                'overrideable': 'false',
                'type': 'string'
            })
            ET.SubElement(dyn_field, 'helpText').text = (
                "Type or paste a SQL prepared statement that is valid for the Update statement. "
                "For more than one statement, separate by semicolon and append a connection property "
                "allowMultiQueries=true to the database url."
            )
            ET.SubElement(dyn_field, 'defaultValue').text = sql_query or ""
            
            opts = ET.SubElement(gen_config, 'Options')
            qopts = ET.SubElement(opts, 'QueryOptions')
            fields_el = ET.SubElement(qopts, 'Fields')
            conn_obj = ET.SubElement(fields_el, 'ConnectorObject', {'name': obj_type_name})
            field_list = ET.SubElement(conn_obj, 'FieldList')
            ET.SubElement(field_list, 'ConnectorField', {'filterable': 'true', 'name': 'Query', 'selectable': 'true', 'selected': 'true', 'sortable': 'true'})
            ET.SubElement(field_list, 'ConnectorField', {'filterable': 'true', 'name': 'Rows Effected', 'selectable': 'true', 'selected': 'true', 'sortable': 'true'})
            ET.SubElement(field_list, 'ConnectorField', {'filterable': 'true', 'name': 'Status', 'selectable': 'true', 'selected': 'true', 'sortable': 'true'})
            ET.SubElement(qopts, 'Inputs')
        elif op_type == "DELETE":
            ET.SubElement(gen_config, 'field', {'id': 'DeleteType', 'type': 'string', 'value': 'Dynamic Delete'})
            ET.SubElement(gen_config, 'field', {'id': 'schemaName', 'type': 'string', 'value': schema_name or ''})
            ET.SubElement(gen_config, 'field', {'id': 'query', 'type': 'string', 'value': sql_query or ""})
            ET.SubElement(gen_config, 'field', {'id': 'joinTransaction', 'type': 'boolean', 'value': 'false'})
            ET.SubElement(gen_config, 'field', {'id': 'CommitOption', 'type': 'string', 'value': 'Commit By Profile'})
            ET.SubElement(gen_config, 'field', {'id': 'batchCount', 'type': 'integer'})
            
        return ET.tostring(comp, encoding='utf-8').decode('utf-8')

    def migrate_map_component(self, map_xml_str, legacy_from_profile_xml, json_from_profile_xml, new_from_profile_id, legacy_to_profile_xml, json_to_profile_xml, new_to_profile_id):
        """
        Rule 2: Fetch profile mappings from Legacy Mapping configurations and map it exactly the same.
        This function rewrites map mappings to target new JSON profile keys for both source and target.
        """
        # Parse legacy source DB profile key->name mapping
        legacy_from_db = ET.fromstring(legacy_from_profile_xml)
        from_elements = legacy_from_db.findall('.//*[@key][@name]')
        from_key_to_name = {el.get('key'): el.get('name') for el in from_elements if el.get('key') and el.get('name')}
        
        # Parse new source JSON profile name->key mapping
        new_from_json = ET.fromstring(json_from_profile_xml)
        new_from_elements = new_from_json.findall('.//*[@key][@name]')
        from_name_to_key = {el.get('name'): el.get('key') for el in new_from_elements if el.get('name') and el.get('key')}
        
        # Parse legacy target DB profile key->name mapping
        legacy_to_db = ET.fromstring(legacy_to_profile_xml)
        to_elements = legacy_to_db.findall('.//*[@key][@name]')
        to_key_to_name = {el.get('key'): el.get('name') for el in to_elements if el.get('key') and el.get('name')}
        
        # Parse new target JSON profile name->key mapping
        new_to_json = ET.fromstring(json_to_profile_xml)
        new_to_elements = new_to_json.findall('.//*[@key][@name]')
        to_name_to_key = {el.get('name'): el.get('key') for el in new_to_elements if el.get('name') and el.get('key')}
        
        # Parse Map XML
        map_root = ET.fromstring(map_xml_str)
        map_el = map_root.find('.//Map')
        if map_el is None:
            raise ValueError("Invalid Map component XML: Map element missing.")
            
        # Set new profile IDs in the Map definition
        map_el.set("fromProfile", new_from_profile_id)
        map_el.set("toProfile", new_to_profile_id)
        
        # Rewrite individual Mappings
        mappings = map_root.findall('.//Mappings/Mapping')
        for mapping in mappings:
            from_key = mapping.get("fromKey")
            to_key = mapping.get("toKey")
            
            # Map source
            if from_key:
                field_name = from_key_to_name.get(from_key)
                if field_name:
                    new_key = from_name_to_key.get(field_name)
                    if new_key:
                        mapping.set("fromKey", new_key)
                        mapping.set("fromKeyPath", f"*[@key='1']/*[@key='2']/*[@key='{new_key}']")
                        
            # Map target
            if to_key:
                field_name = to_key_to_name.get(to_key)
                if field_name:
                    new_key = to_name_to_key.get(field_name)
                    if new_key:
                        mapping.set("toKey", new_key)
                        mapping.set("toKeyPath", f"*[@key='1']/*[@key='2']/*[@key='{new_key}']")
                        
        return ET.tostring(map_root, encoding='utf-8').decode('utf-8')

    def update_map_for_db_profile(self, map_xml_str, legacy_profile_id, new_json_profile_id, legacy_profile_xml, json_profile_xml):
        """
        Dynamically updates a Map component when a legacy DB profile is migrated to a JSON profile.
        Supports DB profile as fromProfile, toProfile, or both, preserving any non-DB profiles.
        """
        ET.register_namespace('bns', 'http://api.platform.boomi.com/')
        ET.register_namespace('xsi', 'http://www.w3.org/2001/XMLSchema-instance')
        map_root = ET.fromstring(map_xml_str)
        map_el = map_root.find('.//Map')
        if map_el is None:
            return map_xml_str
            
        from_prof = map_el.get("fromProfile")
        to_prof = map_el.get("toProfile")
        
        updated = False
        
        # Build field mapping from legacy DB profile (key -> field name)
        legacy_db = ET.fromstring(legacy_profile_xml)
        legacy_elements = legacy_db.findall('.//*[@key][@name]')
        legacy_key_to_name = {el.get('key'): el.get('name') for el in legacy_elements if el.get('key') and el.get('name')}
        
        # Build field mapping from new JSON profile (field name -> key)
        new_json = ET.fromstring(json_profile_xml)
        new_elements = new_json.findall('.//*[@key][@name]')
        new_name_to_key = {el.get('name'): el.get('key') for el in new_elements if el.get('name') and el.get('key')}
        
        # 1. If fromProfile matches legacy DB profile
        if from_prof == legacy_profile_id:
            map_el.set("fromProfile", new_json_profile_id)
            updated = True
            for m in map_root.findall('.//Mappings/Mapping'):
                f_key = m.get("fromKey")
                if f_key:
                    name = legacy_key_to_name.get(f_key)
                    if name and name in new_name_to_key:
                        new_k = new_name_to_key[name]
                        m.set("fromKey", new_k)
                        m.set("fromKeyPath", f"*[@key='1']/*[@key='2']/*[@key='{new_k}']")
                        
        # 2. If toProfile matches legacy DB profile
        if to_prof == legacy_profile_id:
            map_el.set("toProfile", new_json_profile_id)
            updated = True
            for m in map_root.findall('.//Mappings/Mapping'):
                t_key = m.get("toKey")
                if t_key:
                    name = legacy_key_to_name.get(t_key)
                    if name and name in new_name_to_key:
                        new_k = new_name_to_key[name]
                        m.set("toKey", new_k)
                        m.set("toKeyPath", f"*[@key='1']/*[@key='2']/*[@key='{new_k}']")
                        
        if updated:
            return ET.tostring(map_root, encoding='utf-8').decode('utf-8')
        return map_xml_str

    def migrate_process_connector_shapes(self, process_xml_str, legacy_conn_id, new_conn_id, legacy_oper_id, new_oper_id, target_action_type=None):
        """
        Scans process XML to replace legacy database connector references with modern DB V2.
        """
        ET.register_namespace('bns', 'http://api.platform.boomi.com/')
        ET.register_namespace('xsi', 'http://www.w3.org/2001/XMLSchema-instance')
        root = ET.fromstring(process_xml_str)
        shapes = root.findall('.//shapes/shape[@shapetype="connectoraction"]')
        start_shape = root.find('.//shapes/shape[@shapetype="start"]')
        if start_shape is not None:
            shapes.append(start_shape)
        
        count = 0
        for shape in shapes:
            action = shape.find('.//connectoraction')
            if action is not None:
                if action.get("operationId") == legacy_oper_id:
                    action.set("connectorType", "officialboomi-X3979C-dbv2da-prod")
                    action.set("connectionId", new_conn_id)
                    action.set("operationId", new_oper_id)
                    
                    if target_action_type:
                        action.set("actionType", target_action_type)
                    else:
                        curr_act = (action.get("actionType") or "").strip().lower()
                        if curr_act == "send":
                            action.set("actionType", "UPDATE")
                        elif curr_act == "get":
                            action.set("actionType", "GET")
                        elif curr_act == "execute":
                            action.set("actionType", "UPDATE")
                    count += 1
                    
        if count > 0:
            return ET.tostring(root, encoding='utf-8').decode('utf-8')
        return process_xml_str

    def generate_audit_report_excel(self, scan_results, file_path="legacy_db_audit_report.xlsx", max_where_used=50):
        """
        Generates a comprehensive multi-sheet Excel (.xlsx) workbook for the complete legacy DB discovery.
        Sheets: Summary Overview, DB Connections, DB Operations, DB Profiles, Maps (DB Profiles), Document Caches, DB Processes, Folder Inventory.
        """
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # Remove initial sheet

        # Styles
        header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        sub_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
        title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
        bold_font = Font(name="Calibri", size=11, bold=True)
        thin_border = Border(
            left=Side(style='thin', color='D3D3D3'),
            right=Side(style='thin', color='D3D3D3'),
            top=Side(style='thin', color='D3D3D3'),
            bottom=Side(style='thin', color='D3D3D3')
        )

        conns = scan_results.get("connections", [])
        opers = scan_results.get("operations", [])
        profs = scan_results.get("profiles", [])
        maps = scan_results.get("maps", [])
        caches = scan_results.get("caches", [])
        procs = scan_results.get("processes", [])

        # 1. SUMMARY SHEET
        ws_sum = wb.create_sheet(title="Summary Overview")
        ws_sum.views.sheetView[0].showGridLines = True
        ws_sum.append(["BOOMI DATABASE V2 MIGRATION - AUDIT REPORT"])
        ws_sum.cell(row=1, column=1).font = title_font
        ws_sum.append([])
        
        ws_sum.append(["Component Type", "Discovered Count", "Target Migration Action"])
        for col in range(1, 4):
            c = ws_sum.cell(row=3, column=col)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center" if col > 1 else "left")

        summary_rows = [
            ("Database (Legacy) Connections", len(conns), "Migrate to Database V2 Connection (officialboomi-X3979C-dbv2da-prod)"),
            ("Database (Legacy) Operations", len(opers), "Migrate to Database V2 Operation with JSON Request/Response Profiles"),
            ("Database (Legacy) Profiles", len(profs), "Generate corresponding JSON Profiles for DB V2 payload mapping"),
            ("Maps referencing DB Profiles", len(maps), "Update Map shape source/target profiles to JSON Profiles"),
            ("Document Caches using DB Profiles", len(caches), "Reconfigure Document Cache profile definition to modern JSON Profile"),
            ("Processes referencing DB Components", len(procs), "Update Connector Shapes and Map Profiles to DB V2")
        ]
        for row_idx, r in enumerate(summary_rows, start=4):
            ws_sum.append(list(r))
            for col in range(1, 4):
                cell = ws_sum.cell(row=row_idx, column=col)
                cell.border = thin_border
                if col == 2:
                    cell.alignment = Alignment(horizontal="center")
                    cell.font = bold_font

        # Total Legacy Components row
        total_row = len(summary_rows) + 4
        ws_sum.cell(row=total_row, column=1, value="Total Direct Legacy Database Components")
        ws_sum.cell(row=total_row, column=1).font = bold_font
        ws_sum.cell(row=total_row, column=2, value=len(conns) + len(opers) + len(profs))
        ws_sum.cell(row=total_row, column=2).font = bold_font
        ws_sum.cell(row=total_row, column=2).alignment = Alignment(horizontal="center")
        ws_sum.cell(row=total_row, column=3, value="")
        for col in range(1, 4):
            ws_sum.cell(row=total_row, column=col).fill = sub_fill
            ws_sum.cell(row=total_row, column=col).border = thin_border

        # Total Impacted Assets row
        impact_row = total_row + 1
        ws_sum.cell(row=impact_row, column=1, value="Total Discovered & Impacted Artifacts (All Categories)")
        ws_sum.cell(row=impact_row, column=1).font = bold_font
        ws_sum.cell(row=impact_row, column=2, value=len(conns) + len(opers) + len(profs) + len(maps) + len(caches) + len(procs))
        ws_sum.cell(row=impact_row, column=2).font = bold_font
        ws_sum.cell(row=impact_row, column=2).alignment = Alignment(horizontal="center")
        ws_sum.cell(row=impact_row, column=3, value="")
        for col in range(1, 4):
            ws_sum.cell(row=impact_row, column=col).fill = sub_fill
            ws_sum.cell(row=impact_row, column=col).border = thin_border

        for col in ws_sum.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_sum.column_dimensions[col_letter].width = max(max_len + 4, 14)

        # Helper to create component sheets
        def add_component_sheet(title, items):
            ws = wb.create_sheet(title=title)
            ws.views.sheetView[0].showGridLines = True
            headers = ["#", "Component Name", "Component ID", "Version", "Folder Path", "Boomi Type"]
            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=1, column=col_idx)
                c.fill = header_fill
                c.font = header_font
                c.alignment = Alignment(horizontal="center" if col_idx in (1, 3, 4) else "left")

            for idx, item in enumerate(items, start=1):
                row_data = [
                    idx,
                    item.get("name", ""),
                    item.get("componentId", item.get("id", "")),
                    item.get("version", 1),
                    item.get("folderPath", item.get("folderName", "Root")),
                    item.get("type", "")
                ]
                ws.append(row_data)
                row_num = idx + 1
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_num, column=col_idx)
                    cell.border = thin_border
                    if col_idx in (1, 3, 4):
                        cell.alignment = Alignment(horizontal="center")

            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = min(max(max_len + 4, 10), 60)

        # 2. Connections Sheet
        add_component_sheet("DB Connections", conns)

        # 3. Operations Sheet
        add_component_sheet("DB Operations", opers)

        # 4. Profiles Sheet
        add_component_sheet("DB Profiles", profs)

        # 5. Maps Sheet
        add_component_sheet("Maps (DB Profiles)", maps)

        # 6. Caches Sheet
        add_component_sheet("Document Caches", caches)

        # 7. Processes Sheet
        add_component_sheet("DB Processes", procs)

        # 8. Hierarchical Folder Inventory Sheet
        ws_folders = wb.create_sheet(title="Folder Inventory")
        ws_folders.views.sheetView[0].showGridLines = True
        f_headers = [
            "Folder Hierarchy",
            "Folder Type",
            "Folder Full Path",
            "Parent Folder",
            "Total Rollup Items",
            "Direct Items",
            "Subfolder Count",
            "Rollup Processes",
            "Direct Processes",
            "Rollup Maps",
            "Direct Maps",
            "Rollup Caches",
            "Direct Caches",
            "Rollup Conns",
            "Direct Conns",
            "Rollup Ops",
            "Direct Ops",
            "Rollup Profiles",
            "Direct Profiles"
        ]
        ws_folders.append(f_headers)
        for col_idx in range(1, len(f_headers) + 1):
            c = ws_folders.cell(row=1, column=col_idx)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center" if col_idx not in (1, 3) else "left")

        # Build Folder Hierarchy Tree in Python
        class FolderNode:
            def __init__(self, name, full_path, parent_path, depth):
                self.name = name
                self.full_path = full_path
                self.parent_path = parent_path
                self.depth = depth
                self.children = {}
                self.direct = {"conns": 0, "opers": 0, "profs": 0, "maps": 0, "caches": 0, "procs": 0}
                self.direct_count = 0
                self.rollup = {"conns": 0, "opers": 0, "profs": 0, "maps": 0, "caches": 0, "procs": 0}
                self.rollup_count = 0

        folder_tree_root = FolderNode("Root", "", "", 0)
        all_folder_nodes = {}

        root_name = self.get_root_folder_name()

        def get_or_create_node(path_str):
            norm = self._normalize_path(path_str)
            if norm in all_folder_nodes:
                return all_folder_nodes[norm]

            segments = norm.split('/')
            curr = folder_tree_root
            curr_full = ""
            for idx, seg in enumerate(segments):
                parent_full = curr_full
                curr_full = f"{curr_full}/{seg}" if curr_full else seg
                if seg not in curr.children:
                    new_node = FolderNode(seg, curr_full, parent_full or "Root", idx + 1)
                    curr.children[seg] = new_node
                    all_folder_nodes[curr_full] = new_node
                curr = curr.children[seg]
            return curr

        # Place components
        for c in conns:
            node = get_or_create_node(c.get("folderPath") or c.get("folderName") or root_name)
            node.direct["conns"] += 1
        for o in opers:
            node = get_or_create_node(o.get("folderPath") or o.get("folderName") or root_name)
            node.direct["opers"] += 1
        for p in profs:
            node = get_or_create_node(p.get("folderPath") or p.get("folderName") or root_name)
            node.direct["profs"] += 1
        for m in maps:
            node = get_or_create_node(m.get("folderPath") or m.get("folderName") or root_name)
            node.direct["maps"] += 1
        for ca in caches:
            node = get_or_create_node(ca.get("folderPath") or ca.get("folderName") or root_name)
            node.direct["caches"] += 1
        for pr in procs:
            node = get_or_create_node(pr.get("folderPath") or pr.get("folderName") or root_name)
            node.direct["procs"] += 1

        # Post-order rollup calculation
        def calculate_node_rollups(node):
            node.direct_count = sum(node.direct.values())
            node.rollup = dict(node.direct)
            for child in node.children.values():
                calculate_node_rollups(child)
                for k in node.rollup:
                    node.rollup[k] += child.rollup[k]
            node.rollup_count = sum(node.rollup.values())

        calculate_node_rollups(folder_tree_root)

        # Root folders: top-level Main Folders dynamically identified
        effective_roots = list(folder_tree_root.children.values())

        # Output in true hierarchical DFS order (Main Folder -> Subfolders -> Sub-subfolders)
        ordered_entries = []
        def dfs_collect(node, rel_depth):
            ordered_entries.append((node, rel_depth))
            for child_name in sorted(node.children.keys()):
                child_node = node.children[child_name]
                dfs_collect(child_node, rel_depth + 1)

        for root_node in sorted(effective_roots, key=lambda n: n.name):
            dfs_collect(root_node, 1)

        for r_idx, (node, rel_depth) in enumerate(ordered_entries, start=2):
            folder_type = "Main Folder" if rel_depth == 1 else f"Subfolder (Level {rel_depth - 1})"
            indent_prefix = "    " * (rel_depth - 1) + ("└─ " if rel_depth > 1 else "")
            display_name = f"{indent_prefix}{node.name}"
            row_data = [
                display_name,
                folder_type,
                node.full_path,
                node.parent_path,
                node.rollup_count,
                node.direct_count,
                len(node.children),
                node.rollup["procs"],
                node.direct["procs"],
                node.rollup["maps"],
                node.direct["maps"],
                node.rollup["caches"],
                node.direct["caches"],
                node.rollup["conns"],
                node.direct["conns"],
                node.rollup["opers"],
                node.direct["opers"],
                node.rollup["profs"],
                node.direct["profs"]
            ]
            ws_folders.append(row_data)
            for c_idx in range(1, len(f_headers) + 1):
                cell = ws_folders.cell(row=r_idx, column=c_idx)
                cell.border = thin_border
                if c_idx in (1, 2, 5, 6):
                    cell.font = bold_font
                if c_idx in (2, 4, 5, 6, 7) or c_idx > 7:
                    cell.alignment = Alignment(horizontal="center")

        for col in ws_folders.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws_folders.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 70)

        wb.save(file_path)
        return file_path

    def generate_audit_report_md(self, scan_results):
        """
        Generates markdown audit report of discovered components.
        """
        conns = scan_results.get("connections", [])
        opers = scan_results.get("operations", [])
        profs = scan_results.get("profiles", [])
        maps = scan_results.get("maps", [])
        caches = scan_results.get("caches", [])
        procs = scan_results.get("processes", [])

        md = []
        md.append("# Boomi Database (Legacy) to Database V2 Migration Audit Report\n")
        md.append("## Executive Summary\n")
        md.append(f"- **Total DB Connections (Legacy)**: {len(conns)}")
        md.append(f"- **Total DB Operations (Legacy)**: {len(opers)}")
        md.append(f"- **Total DB Profiles (Legacy)**: {len(profs)}")
        md.append(f"- **Maps referencing DB Profiles**: {len(maps)}")
        md.append(f"- **Document Caches referencing DB Profiles**: {len(caches)}")
        md.append(f"- **Discovered DB Processes**: {len(procs)}")
        md.append(f"- **Total Impacted Artifacts**: {len(conns) + len(opers) + len(profs) + len(maps) + len(caches) + len(procs)}\n")
        
        md.append("## Discovered Database Connections\n")
        md.append("| Name | Component ID | Folder Path |")
        md.append("|---|---|---|")
        for c in conns[:100]:
            md.append(f"| {c.get('name')} | `{c.get('componentId')}` | {c.get('folderPath', 'Root')} |")
        if len(conns) > 100:
            md.append(f"| *...and {len(conns)-100} more connections* | | |")
            
        md.append("\n## Discovered Database Operations\n")
        md.append("| Name | Component ID | Folder Path |")
        md.append("|---|---|---|")
        for o in opers[:100]:
            md.append(f"| {o.get('name')} | `{o.get('componentId')}` | {o.get('folderPath', 'Root')} |")
        if len(opers) > 100:
            md.append(f"| *...and {len(opers)-100} more operations* | | |")

        md.append("\n## Discovered Database Profiles\n")
        md.append("| Name | Component ID | Folder Path |")
        md.append("|---|---|---|")
        for p in profs[:100]:
            md.append(f"| {p.get('name')} | `{p.get('componentId')}` | {p.get('folderPath', 'Root')} |")
        if len(profs) > 100:
            md.append(f"| *...and {len(profs)-100} more profiles* | | |")

        md.append("\n## Discovered Maps using DB Profiles\n")
        md.append("| Name | Component ID | Folder Path |")
        md.append("|---|---|---|")
        for m in maps[:100]:
            md.append(f"| {m.get('name')} | `{m.get('componentId')}` | {m.get('folderPath', 'Root')} |")
        if len(maps) > 100:
            md.append(f"| *...and {len(maps)-100} more maps* | | |")

        md.append("\n## Discovered Document Caches using DB Profiles\n")
        md.append("| Name | Component ID | Folder Path |")
        md.append("|---|---|---|")
        for ca in caches[:100]:
            md.append(f"| {ca.get('name')} | `{ca.get('componentId')}` | {ca.get('folderPath', 'Root')} |")
        if len(caches) > 100:
            md.append(f"| *...and {len(caches)-100} more caches* | | |")

        md.append("\n## Discovered Processes Referencing DB Components\n")
        md.append("| Name | Component ID | Folder Path |")
        md.append("|---|---|---|")
        for pr in procs[:100]:
            md.append(f"| {pr.get('name')} | `{pr.get('componentId')}` | {pr.get('folderPath', 'Root')} |")
        if len(procs) > 100:
            md.append(f"| *...and {len(procs)-100} more processes* | | |")

        return "\n".join(md)
