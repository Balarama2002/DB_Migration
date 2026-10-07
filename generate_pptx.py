import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def hex_to_rgb(hex_str):
    hex_str = hex_str.lstrip('#')
    return RGBColor(*(int(hex_str[i:i+2], 16) for i in (0, 2, 4)))

def create_deck():
    prs = Presentation()
    # 13.333 x 8.0 inches (widescreen presentation, matches 1300x780 aspect ratio)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(8.0)
    
    # Scale factors: SVG is 1300 x 780
    sx = 13.333 / 1300.0
    sy = 8.0 / 780.0
    
    def to_x(x): return Inches(x * sx)
    def to_y(y): return Inches(y * sy)
    def to_w(w): return Inches(w * sx)
    def to_h(h): return Inches(h * sy)
    
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    
    # Background light gray
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = hex_to_rgb('fafbfc')
    bg.line.fill.background()
    
    # 1. Outer Container (Green Border)
    outer = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(25), to_y(25), to_w(1250), to_h(730))
    outer.fill.solid()
    outer.fill.fore_color.rgb = hex_to_rgb('ffffff')
    outer.line.color.rgb = hex_to_rgb('16a34a')
    outer.line.width = Pt(2.5)
    
    # Title & Subtitle (Top Left)
    title_box = slide.shapes.add_textbox(to_x(50), to_y(38), to_w(420), to_h(50))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "Boomi Database V2 Migration Agent"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('16a34a')
    
    p2 = tf.add_paragraph()
    p2.text = "End-to-End Modernization Suite: Legacy Database → Database V2 (officialboomi-X3979C-dbv2da-prod)"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = hex_to_rgb('64748b')
    
    # Migration Hub UI (Top Center)
    ui_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(490), to_y(32), to_w(230), to_h(58))
    ui_card.fill.solid()
    ui_card.fill.fore_color.rgb = hex_to_rgb('f0fdf4')
    ui_card.line.color.rgb = hex_to_rgb('22c55e')
    ui_card.line.width = Pt(1.5)
    tf_ui = ui_card.text_frame
    tf_ui.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf_ui.paragraphs[0]
    p.text = "🖥️ Migration Hub UI"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0f172a')
    p2 = tf_ui.add_paragraph()
    p2.text = "Mode A (Single) | Mode B (Bulk Folder)"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = hex_to_rgb('15803d')
    
    # Flow Legend (Top Right)
    legend_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(940), to_y(32), to_w(315), to_h(62))
    legend_box.fill.solid()
    legend_box.fill.fore_color.rgb = hex_to_rgb('f8fafc')
    legend_box.line.color.rgb = hex_to_rgb('cbd5e1')
    legend_box.line.width = Pt(1)
    tf_leg = legend_box.text_frame
    tf_leg.word_wrap = True
    tf_leg.margin_left = Inches(0.1)
    tf_leg.margin_top = Inches(0.04)
    p = tf_leg.paragraphs[0]
    p.text = "🔴 Primary Flow    🔵 Discovery Flow"
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('334155')
    p2 = tf_leg.add_paragraph()
    p2.text = "🟠 Transformation Flow    🟢 Deployment Response"
    p2.font.size = Pt(8.5)
    p2.font.bold = True
    p2.font.color.rgb = hex_to_rgb('334155')
    
    # 2. Main Inner Box (Red Border: AI-Engine)
    ai_engine = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(45), to_y(105), to_w(1210), to_h(635))
    ai_engine.fill.solid()
    ai_engine.fill.fore_color.rgb = hex_to_rgb('ffffff')
    ai_engine.line.color.rgb = hex_to_rgb('ef4444')
    ai_engine.line.width = Pt(2)
    
    # AI-Engine Header Badge
    ai_header = slide.shapes.add_textbox(to_x(65), to_y(112), to_w(380), to_h(40))
    tf_aih = ai_header.text_frame
    tf_aih.word_wrap = True
    tf_aih.margin_left = tf_aih.margin_top = tf_aih.margin_right = tf_aih.margin_bottom = 0
    p = tf_aih.paragraphs[0]
    p.text = "🤖 AI-Engine"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('2563eb')
    p2 = tf_aih.add_paragraph()
    p2.text = "MigrationTools Orchestration & API Integration"
    p2.font.size = Pt(9.5)
    p2.font.color.rgb = hex_to_rgb('64748b')
    
    # 3. Left Column: Boomi Runtime Engine
    runtime_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(175), to_y(215), to_w(160), to_h(120))
    runtime_card.fill.solid()
    runtime_card.fill.fore_color.rgb = hex_to_rgb('f8fafc')
    runtime_card.line.color.rgb = hex_to_rgb('334155')
    runtime_card.line.width = Pt(2)
    tf_rt = runtime_card.text_frame
    tf_rt.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf_rt.paragraphs[0]
    p.text = "⚛️ Boomi Runtime Engine"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0f172a')
    p2 = tf_rt.add_paragraph()
    p2.text = "FastAPI Migration Hub\nSSE Live Event Stream"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = hex_to_rgb('0284c7')
    
    # Metadata Exchange Badge
    meta_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(185), to_y(395), to_w(140), to_h(55))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = hex_to_rgb('f1f5f9')
    meta_card.line.color.rgb = hex_to_rgb('0f172a')
    meta_card.line.width = Pt(1.5)
    tf_mc = meta_card.text_frame
    tf_mc.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf_mc.paragraphs[0]
    p.text = "↕️ Metadata Exchange"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0f172a')
    p2 = tf_mc.add_paragraph()
    p2.text = "Component API"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8)
    p2.font.color.rgb = hex_to_rgb('64748b')
    
    # Boomi Enterprise Platform Cloud Card
    cloud_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(140), to_y(510), to_w(230), to_h(125))
    cloud_card.fill.solid()
    cloud_card.fill.fore_color.rgb = hex_to_rgb('f0f9ff')
    cloud_card.line.color.rgb = hex_to_rgb('0284c7')
    cloud_card.line.width = Pt(2)
    tf_cc = cloud_card.text_frame
    tf_cc.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf_cc.paragraphs[0]
    p.text = "☁️ Boomi Enterprise Platform"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0369a1')
    p2 = tf_cc.add_paragraph()
    p2.text = "AtomSphere Integration APIs\nBranch: dbv2_merging\n🛡️ Main Branch Guard: Active"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = hex_to_rgb('0f172a')
    
    # 4. Top Right Box: Knowledge-Base (Blue Border)
    kb_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(460), to_y(120), to_w(775), to_h(225))
    kb_box.fill.solid()
    kb_box.fill.fore_color.rgb = hex_to_rgb('ffffff')
    kb_box.line.color.rgb = hex_to_rgb('2563eb')
    kb_box.line.width = Pt(2)
    
    # KB Header
    kb_hdr = slide.shapes.add_textbox(to_x(475), to_y(126), to_w(740), to_h(30))
    tf_kbh = kb_hdr.text_frame
    p = tf_kbh.paragraphs[0]
    p.text = "Knowledge-Base — Hierarchical Subfolder Discovery & Where-Used Dependency Engine"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('2563eb')
    
    # KB Stage 1: Subfolder Audit
    kb_s1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(485), to_y(165), to_w(105), to_h(125))
    kb_s1.fill.solid()
    kb_s1.fill.fore_color.rgb = hex_to_rgb('eff6ff')
    kb_s1.line.color.rgb = hex_to_rgb('3b82f6')
    kb_s1.line.width = Pt(1.5)
    tf = kb_s1.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "📁\nSubfolder Audit"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0f172a')
    p2 = tf.add_paragraph()
    p2.text = "Recursive parentId\nTree & Rollups"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8)
    p2.font.color.rgb = hex_to_rgb('64748b')
    
    # KB Stage 2: DB Discovery
    kb_s2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(605), to_y(165), to_w(105), to_h(125))
    kb_s2.fill.solid()
    kb_s2.fill.fore_color.rgb = hex_to_rgb('eff6ff')
    kb_s2.line.color.rgb = hex_to_rgb('3b82f6')
    kb_s2.line.width = Pt(1.5)
    tf = kb_s2.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "🔍\nDB Discovery"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0f172a')
    p2 = tf.add_paragraph()
    p2.text = "Legacy Conns/Opers\n& profile.db"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8)
    p2.font.color.rgb = hex_to_rgb('64748b')
    
    # KB Stage 3: Where-Used Graph
    kb_s3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(725), to_y(165), to_w(110), to_h(125))
    kb_s3.fill.solid()
    kb_s3.fill.fore_color.rgb = hex_to_rgb('eff6ff')
    kb_s3.line.color.rgb = hex_to_rgb('3b82f6')
    kb_s3.line.width = Pt(1.5)
    tf = kb_s3.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "🔗\nWhere-Used"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0f172a')
    p2 = tf.add_paragraph()
    p2.text = "ComponentReference\nMaps/Caches/Processes"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(8)
    p2.font.color.rgb = hex_to_rgb('64748b')
    
    # Center DB Cache
    kb_db = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(850), to_y(180), to_w(85), to_h(95))
    kb_db.fill.solid()
    kb_db.fill.fore_color.rgb = hex_to_rgb('dbeafe')
    kb_db.line.color.rgb = hex_to_rgb('2563eb')
    kb_db.line.width = Pt(2)
    tf = kb_db.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "🗄️\nMetadata DB\nCache"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('1d4ed8')
    
    # KB Stage 4: Non-DB Filter
    kb_s4 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(950), to_y(165), to_w(85), to_h(125))
    kb_s4.fill.solid()
    kb_s4.fill.fore_color.rgb = hex_to_rgb('eff6ff')
    kb_s4.line.color.rgb = hex_to_rgb('3b82f6')
    kb_s4.line.width = Pt(1.5)
    tf = kb_s4.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "🛡️\nNon-DB Filter"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0f172a')
    p2 = tf.add_paragraph()
    p2.text = "Preserve Disk/HTTP\n& SFTP"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(7.5)
    p2.font.color.rgb = hex_to_rgb('64748b')
    
    # KB Stage 5: Action Classify
    kb_s5 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(1045), to_y(165), to_w(90), to_h(125))
    kb_s5.fill.solid()
    kb_s5.fill.fore_color.rgb = hex_to_rgb('eff6ff')
    kb_s5.line.color.rgb = hex_to_rgb('3b82f6')
    kb_s5.line.width = Pt(1.5)
    tf = kb_s5.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "⚖️\nAction Classify"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0f172a')
    p2 = tf.add_paragraph()
    p2.text = "Dynamic Update/Insert\nDelete / Get"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(7.5)
    p2.font.color.rgb = hex_to_rgb('64748b')
    
    # KB Stage 6: Scope Inventory
    kb_s6 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(1145), to_y(165), to_w(80), to_h(125))
    kb_s6.fill.solid()
    kb_s6.fill.fore_color.rgb = hex_to_rgb('eff6ff')
    kb_s6.line.color.rgb = hex_to_rgb('3b82f6')
    kb_s6.line.width = Pt(1.5)
    tf = kb_s6.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "📑\nScope Inventory"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(9)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('0f172a')
    p2 = tf.add_paragraph()
    p2.text = "Feature Branch\nComponent XMLs"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(7.5)
    p2.font.color.rgb = hex_to_rgb('64748b')
    
    # KB Footer Note
    kb_ftr = slide.shapes.add_textbox(to_x(480), to_y(305), to_w(740), to_h(25))
    tf = kb_ftr.text_frame
    p = tf.paragraphs[0]
    p.text = "• Discovers all 6 component tiers: Connections, Operations, Profiles, Maps, Caches, and DB Processes"
    p.font.size = Pt(9.5)
    p.font.color.rgb = hex_to_rgb('2563eb')
    
    # 5. Bottom Right Box: AI-Reasoning (Orange Border)
    ar_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(460), to_y(360), to_w(775), to_h(365))
    ar_box.fill.solid()
    ar_box.fill.fore_color.rgb = hex_to_rgb('ffffff')
    ar_box.line.color.rgb = hex_to_rgb('f97316')
    ar_box.line.width = Pt(2)
    
    # AR Header
    ar_hdr = slide.shapes.add_textbox(to_x(475), to_y(368), to_w(740), to_h(30))
    tf_arh = ar_hdr.text_frame
    p = tf_arh.paragraphs[0]
    p.text = "AI-Reasoning — Database V2 Modernization Cycle (officialboomi-X3979C-dbv2da-prod)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('f97316')
    
    # Central Core Engine Chip
    center_chip = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(795), to_y(500), to_w(105), to_h(90))
    center_chip.fill.solid()
    center_chip.fill.fore_color.rgb = hex_to_rgb('ffedd5')
    center_chip.line.color.rgb = hex_to_rgb('ea580c')
    center_chip.line.width = Pt(2.5)
    tf = center_chip.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "⚡\nDB V2\nENGINE"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('c2410c')
    
    # The 6 Concrete Phases arranged in a 3-column x 2-row clean modular grid
    phases = [
        ("🔌 1. Connection Modernization", "Convert Legacy JDBC Settings\nofficialboomi-X3979C-dbv2da-prod\nHost, Port, DB, User credentials", 485, 415),
        ("📄 2. Profile → JSON Synthesis", "profile.db → JSON Profiles\nJSONRootValue > Object > Entries\nNullable & DataType mapping", 745, 415),
        ("⚙️ 3. Operation Assembly", "GenericOperationConfig & objectTypeId\nLink Req & Resp JSON Profiles\nCustom SQL Query / SP configuration", 1005, 415),
        ("🔄 6. Process Shape Re-Wiring", "Rewire <connectoraction> Shapes\nSet dbv2da-prod action type\nPsiog Peer Review Standards", 485, 595),
        ("🔀 5. Map Shape Re-Wiring", "Swap fromProfile / toProfile to JSON\nRemap field keys & key paths\nRetain user functions & lookups", 745, 595),
        ("🧠 4. Dynamic SQL Synthesis", "Prepared Statements & WHERE Logic\nConditions: =, <>, >, <, LIKE\nParameter binding & Type casting", 1005, 595)
    ]
    
    for title, desc, px, py in phases:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(px), to_y(py), to_w(235), to_h(110))
        card.fill.solid()
        card.fill.fore_color.rgb = hex_to_rgb('fff7ed')
        card.line.color.rgb = hex_to_rgb('fb923c')
        card.line.width = Pt(1.5)
        tf = card.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = hex_to_rgb('9a3412')
        p2 = tf.add_paragraph()
        p2.text = desc
        p2.alignment = PP_ALIGN.CENTER
        p2.font.size = Pt(8)
        p2.font.color.rgb = hex_to_rgb('64748b')
        
    # Flow Callout Badges
    callout1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(345), to_y(175), to_w(105), to_h(40))
    callout1.fill.solid()
    callout1.fill.fore_color.rgb = hex_to_rgb('fee2e2')
    callout1.line.color.rgb = hex_to_rgb('dc2626')
    callout1.line.width = Pt(1.5)
    tf = callout1.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "🛡️ Feed to KB\n(Firewall Guard)"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('991b1b')

    callout2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(345), to_y(285), to_w(105), to_h(35))
    callout2.fill.solid()
    callout2.fill.fore_color.rgb = hex_to_rgb('fee2e2')
    callout2.line.color.rgb = hex_to_rgb('dc2626')
    callout2.line.width = Pt(1.5)
    tf = callout2.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "⬅️ Fetch from KB"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('991b1b')

    callout3 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(345), to_y(445), to_w(105), to_h(35))
    callout3.fill.solid()
    callout3.fill.fore_color.rgb = hex_to_rgb('ffedd5')
    callout3.line.color.rgb = hex_to_rgb('ea580c')
    callout3.line.width = Pt(1.5)
    tf = callout3.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "➡️ Invoke AI"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('c2410c')

    callout4 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, to_x(345), to_y(565), to_w(105), to_h(35))
    callout4.fill.solid()
    callout4.fill.fore_color.rgb = hex_to_rgb('dcfce7')
    callout4.line.color.rgb = hex_to_rgb('16a34a')
    callout4.line.width = Pt(1.5)
    tf = callout4.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = "⬅️ AI Response"
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(8.5)
    p.font.bold = True
    p.font.color.rgb = hex_to_rgb('15803d')

    output_path = "db_migration_architecture.pptx"
    prs.save(output_path)
    print(f"Presentation saved to {output_path}")

if __name__ == "__main__":
    create_deck()
